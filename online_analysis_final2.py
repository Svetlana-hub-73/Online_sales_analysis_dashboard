import pandas as pd
import seaborn as sns
import matplotlib.pyplot as plt
import streamlit as st
from pptx import Presentation
from pptx.util import Inches
import os

# Настройка визуального стиля
sns.set(style="whitegrid")

# ───────── Загрузка данных ───────── #
df = pd.read_csv("online_store_data.csv")

# ───────── Предобработка ───────── #
df['category'] = df['category'].astype(str).str.strip().str.lower()
df['product_name'] = df['product_name'].astype(str).str.lower()
df['brand'] = df['brand'].astype(str).str.strip()
df['price'] = pd.to_numeric(df['price'], errors='coerce')
df['quantity_sold'] = pd.to_numeric(df['quantity_sold'], errors='coerce')
df['rating'] = df['rating'].astype(str).str.extract(r'([\d.]+)').astype(float)
df['num_of_ratings'] = pd.to_numeric(df['num_of_ratings'], errors='coerce')

# ───────── Streamlit UI ───────── #
st.title("🛒 Анализ онлайн-продаж с визуализациями и презентацией")

# Фильтры
category_filter = st.selectbox("Выберите категорию", options=["все"] + sorted(df['category'].unique()))
brand_filter = st.selectbox("Выберите бренд", options=["все"] + sorted(df['brand'].dropna().unique()))

filtered_df = df.copy()
if category_filter != "все":
    filtered_df = filtered_df[filtered_df['category'] == category_filter]
if brand_filter != "все":
    filtered_df = filtered_df[filtered_df['brand'] == brand_filter]

st.write(f"📦 Найдено {len(filtered_df)} товаров по выбранным фильтрам")

# ───────── Анализ: TV ───────── #
tv_df = filtered_df[filtered_df['category'].str.contains('tv') | filtered_df['product_name'].str.contains('tv')]
if not tv_df.empty:
    tv_rating_diff = tv_df['rating'].max() - tv_df['rating'].min()
    st.subheader("1. Разница рейтинга среди телевизоров")
    st.write(f"Разница: {tv_rating_diff:.2f}")
    fig1, ax1 = plt.subplots()
    sns.histplot(tv_df['rating'], bins=10, kde=True, ax=ax1)
    ax1.set_title("Распределение рейтингов ТВ")
    st.pyplot(fig1)

# ───────── Анализ: Смартфоны ───────── #
smartphones_df = filtered_df[filtered_df['category'].str.contains('smartphone')]
if not smartphones_df.empty:
    q1 = smartphones_df['price'].quantile(0.25)
    q3 = smartphones_df['price'].quantile(0.75)
    iqr = q3 - q1
    st.subheader("2. Межквартильный диапазон цен смартфонов")
    st.write(f"IQR: {iqr:.1f}")
    fig2, ax2 = plt.subplots()
    sns.boxplot(x=smartphones_df['price'], ax=ax2)
    ax2.set_title("Boxplot цен смартфонов")
    st.pyplot(fig2)

# ───────── Анализ: Стабильные бренды ───────── #
stable_brands_df = filtered_df.dropna(subset=['brand', 'rating'])
brand_std = stable_brands_df.groupby('brand')['rating'].std().dropna().sort_values().head(5)

st.subheader("3. ТОП-5 брендов с самыми стабильными рейтингами")
st.dataframe(brand_std)

fig3, ax3 = plt.subplots()
brand_std.plot(kind='barh', ax=ax3, color='skyblue')
ax3.set_title("Минимальное отклонение рейтингов по брендам")
st.pyplot(fig3)

# ───────── Анализ: Продажи по квартилям ───────── #
filtered_df['rating_quartile'] = pd.qcut(filtered_df['num_of_ratings'], 4, labels=[
    '1-й квартиль', '2-й квартиль', '3-й квартиль', '4-й квартиль'
])
quartile_sales = filtered_df.groupby('rating_quartile', observed=True)['quantity_sold'].sum()

st.subheader("4. Продажи по квартилям количества оценок")
st.dataframe(quartile_sales)

fig4, ax4 = plt.subplots()
sns.barplot(x=quartile_sales.index, y=quartile_sales.values,
            hue=quartile_sales.index, palette='viridis', legend=False, ax=ax4)
ax4.set_title("Продажи по квартилям")
st.pyplot(fig4)

# ───────── Корреляция ───────── #
st.subheader("5. Корреляция между price, rating и quantity_sold")
corr = filtered_df[['price', 'rating', 'quantity_sold']].corr()
st.dataframe(corr)

fig5, ax5 = plt.subplots()
sns.heatmap(corr, annot=True, cmap="coolwarm", fmt=".2f", ax=ax5)
ax5.set_title("Матрица корреляции")
st.pyplot(fig5)

# ───────── Генерация презентации ───────── #
prs = Presentation()
slide_title = prs.slides.add_slide(prs.slide_layouts[0])
slide_title.shapes.title.text = "📈 Анализ данных онлайн-магазина"
slide_title.placeholders[1].text = f"Категория: {category_filter}, Бренд: {brand_filter}"

def add_slide(title, content):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    slide.placeholders[1].text = content

if not tv_df.empty:
    add_slide("1. Рейтинг ТВ", f"Разница рейтингов: {tv_rating_diff:.2f}")
if not smartphones_df.empty:
    add_slide("2. Смартфоны", f"IQR цен: {iqr:.1f}")
add_slide("3. ТОП Бренды", brand_std.to_string())
add_slide("4. Квартильные продажи", quartile_sales.to_string())
add_slide("5. Корреляция", corr.to_string())

pptx_path = "presentation_output.pptx"
prs.save(pptx_path)
st.download_button("📥 Скачать презентацию", data=open(pptx_path, "rb"), file_name="online_analysis.pptx")

# ───────── Excel отчёт ───────── #
excel_path = "filtered_data.xlsx"
filtered_df.to_excel(excel_path, index=False)
st.download_button("📥 Скачать Excel отчёт", data=open(excel_path, "rb"), file_name="filtered_data.xlsx")

