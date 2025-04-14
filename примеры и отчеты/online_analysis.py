import pandas as pd
import seaborn as sns
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches
import streamlit as st

# Загрузка данных
df = pd.read_csv("online_store_data.csv")  # замени на путь к своему CSV

# Очистка и преобразование
df['category'] = df['category'].astype(str).str.strip().str.lower()
df['product_name'] = df['product_name'].astype(str).str.lower()
df['brand'] = df['brand'].astype(str).str.strip()

# Преобразуем нужные столбцы в числовой тип
df['price'] = pd.to_numeric(df['price'], errors='coerce')
df['quantity_sold'] = pd.to_numeric(df['quantity_sold'], errors='coerce')
df['rating'] = df['rating'].str.extract(r'([\d.]+)').astype(float)
df['num_of_ratings'] = pd.to_numeric(df['num_of_ratings'], errors='coerce')

# ───────── 1. Разница рейтинга среди телевизоров ───────── #
tv_df = df[df['category'].str.contains('tv') | df['product_name'].str.contains('tv')]
if not tv_df.empty:
    tv_rating_diff = tv_df['rating'].max() - tv_df['rating'].min()
    print(f"1. Разница рейтинга среди телевизоров: {tv_rating_diff:.2f}")
else:
    print("1. Телевизоры не найдены в данных")

# ───────── 2. Межквартильный диапазон цен для смартфонов ───────── #
smartphones_df = df[df['category'].str.contains('smartphone')]
if not smartphones_df.empty:
    q1 = smartphones_df['price'].quantile(0.25)
    q3 = smartphones_df['price'].quantile(0.75)
    iqr = q3 - q1
    print(f"\n2. Межквартильный диапазон цен для смартфонов: {iqr:.1f}")
else:
    print("\n2. Смартфоны не найдены")

# ───────── 3. Бренды с наиболее стабильными рейтингами ───────── #
# Убираем бренды с NaN
stable_brands_df = df.dropna(subset=['brand', 'rating'])
brand_std = stable_brands_df.groupby('brand')['rating'].std().dropna().sort_values().head(5)

print("\n3. Бренды с наиболее стабильными рейтингами:")
print(brand_std)

# ───────── 4. Продажи по квартилям количества оценок ───────── #
df['rating_quartile'] = pd.qcut(df['num_of_ratings'], 4, labels=['1-й квартиль', '2-й квартиль', '3-й квартиль', '4-й квартиль'])
quartile_sales = df.groupby('rating_quartile', observed=True)['quantity_sold'].sum()

print("\n4. Продажи по квартилям количества оценок:")
print(quartile_sales)

# ───────── 1. Гистограмма рейтингов телевизоров ───────── #
if not tv_df.empty:
    plt.figure(figsize=(10, 6))
    sns.histplot(tv_df['rating'], bins=10, kde=True, color='blue')
    plt.title("Распределение рейтингов среди телевизоров", fontsize=15)
    plt.xlabel("Рейтинг", fontsize=12)
    plt.ylabel("Частота", fontsize=12)
    plt.savefig("tv_ratings_histogram.png")  # Сохраняем график
    plt.show()

# ───────── 2. Боксплот цен для смартфонов ───────── #
if not smartphones_df.empty:
    plt.figure(figsize=(10, 6))
    sns.boxplot(x=smartphones_df['price'], color='orange')
    plt.title("Распределение цен среди смартфонов", fontsize=15)
    plt.xlabel("Цена", fontsize=12)
    plt.savefig("smartphone_price_boxplot.png")  # Сохраняем график
    plt.show()

# ───────── 3. Диаграмма с точками для стабильных брендов ───────── #
if not brand_std.empty:
    plt.figure(figsize=(10, 6))
    sns.scatterplot(x=brand_std.index, y=brand_std.values, color='green')
    plt.title("Стабильность рейтингов брендов", fontsize=15)
    plt.xlabel("Бренд", fontsize=12)
    plt.ylabel("Стандартное отклонение рейтинга", fontsize=12)
    plt.xticks(rotation=45)
    plt.savefig("brand_stability_scatter.png")  # Сохраняем график
    plt.show()

# ───────── 4. Столбчатая диаграмма продаж по квартилям количества оценок ───────── #
plt.figure(figsize=(10, 6))
sns.barplot(x=quartile_sales.index, y=quartile_sales.values, hue=quartile_sales.index, palette='viridis', legend=False)

plt.title("Продажи по квартилям количества оценок", fontsize=15)
plt.xlabel("Квартиль количества оценок", fontsize=12)
plt.ylabel("Общее количество продаж", fontsize=12)
plt.savefig("sales_by_rating_quartile.png")  # Сохраняем график
plt.show()

# ───────── 5. Создание презентации с результатами анализа ───────── #
# Создание презентации
prs = Presentation()

# Слайд 1: Заголовок
slide_1 = prs.slides.add_slide(prs.slide_layouts[0])
title = slide_1.shapes.title
subtitle = slide_1.placeholders[1]
title.text = "Анализ данных онлайн-магазина"
subtitle.text = "Результаты анализа продаж и рейтингов товаров"

# Слайд 2: Гистограмма рейтингов телевизоров
slide_2 = prs.slides.add_slide(prs.slide_layouts[5])
title = slide_2.shapes.title
title.text = "Распределение рейтингов среди телевизоров"
slide_2.shapes.add_picture("tv_ratings_histogram.png", Inches(1), Inches(1.5), height=Inches(3.5))

# Слайд 3: Боксплот цен для смартфонов
slide_3 = prs.slides.add_slide(prs.slide_layouts[5])
title = slide_3.shapes.title
title.text = "Распределение цен среди смартфонов"
slide_3.shapes.add_picture("smartphone_price_boxplot.png", Inches(1), Inches(1.5), height=Inches(3.5))

# Слайд 4: Диаграмма с точками для стабильных брендов
slide_4 = prs.slides.add_slide(prs.slide_layouts[5])
title = slide_4.shapes.title
title.text = "Стабильность рейтингов брендов"
slide_4.shapes.add_picture("brand_stability_scatter.png", Inches(1), Inches(1.5), height=Inches(3.5))

# Слайд 5: Столбчатая диаграмма продаж по квартилям количества оценок
slide_5 = prs.slides.add_slide(prs.slide_layouts[5])
title = slide_5.shapes.title
title.text = "Продажи по квартилям количества оценок"
slide_5.shapes.add_picture("sales_by_rating_quartile.png", Inches(1), Inches(1.5), height=Inches(3.5))

# Сохранение презентации
prs.save('online_store_analysis.pptx')

# ───────── 6. Streamlit интерфейс ───────── #

st.title("Анализ данных онлайн-магазина")

st.header("Гистограмма рейтингов телевизоров")
st.image("tv_ratings_histogram.png", caption="Распределение рейтингов среди телевизоров")

st.header("Боксплот цен для смартфонов")
st.image("smartphone_price_boxplot.png", caption="Распределение цен среди смартфонов")

st.header("Стабильность рейтингов брендов")
st.image("brand_stability_scatter.png", caption="Стабильность рейтингов брендов")

st.header("Продажи по квартилям количества оценок")
st.image("sales_by_rating_quartile.png", caption="Продажи по квартилям количества оценок")

st.download_button(
    label="Скачать презентацию",
    data=open("online_store_analysis.pptx", "rb").read(),
    file_name="online_store_analysis.pptx",
    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
)
